"""wiki・BRILLI@NT・シャニマスギャラリーの情報で総覧 Excel を更新する。

使い方:
  python scripts/update_excel.py 総覧.xlsx --wiki work/wiki.json --brilliant work/brilliant.json \
      --gamedbs work/gamedbs.json -o work/総覧_new.xlsx --report work/report.json --plan work/sprite_plan.json

やること（ユーザーが書いた必読度・メモ・行は触らない）:
  0. wiki でコミュ数が0のサポートカードの行を外す（必読度・メモが書いてある行は残して報告する）
  1. wiki にあって総覧にない新しいカードを、日付順の位置に挿入する
     （プロデュースは全部、サポートはコミュ数が1以上のものだけ。最新の実装日の30日前以降が対象）
  2. カード行の区分（H列）を wiki のガシャ種別に合わせる。wiki にない区分は BRILLI@NT から補う
  3. カードページ（I列）が空のカード行に、BRILLI@NT → シャニマスギャラリーの順で URL を入れる
  4. 画像がまだないカード（I列のURLが sprites/registry.json にない）を sprite_plan.json に書き出す
  5. 報告用に、追加・除外・区分変更・総覧にないイベント・一覧に未反映のピックアップを report.json に書く
"""
import argparse
import datetime
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from common import (BRILLIANT_ONLY_COLLS, CARD_KINDS, COL, EVENT, GACHA_TO_COLL, KIND_OF, P, S,
                    best_match, dump_json, image_of, iso, load_json, norm, people_of, split_tags, to_date)
from xlsx_util import delete_row, get_link, insert_rows, set_link

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NEW_CARD_WINDOW_DAYS = 30


def read_rows(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        kind = ws.cell(r, COL["kind"]).value
        name = ws.cell(r, COL["name"]).value
        if not kind and not name:
            continue
        rows.append(dict(r=r, kind=kind or "", tags=split_tags(ws.cell(r, COL["tags"]).value),
                         name="" if name is None else str(name), date=to_date(ws.cell(r, COL["date"]).value),
                         coll=ws.cell(r, COL["coll"]).value or "", link=get_link(ws, r)))
    return rows


def find_card(rows, kind, idol, name):
    return best_match(name, [x for x in rows if x["kind"] == kind and idol in x["tags"]], lambda x: x["name"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--wiki", required=True)
    ap.add_argument("--brilliant", required=True)
    ap.add_argument("--gamedbs", required=True)
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--report", required=True)
    ap.add_argument("--plan", required=True)
    a = ap.parse_args()

    wiki = load_json(a.wiki)
    brilliant = load_json(a.brilliant)
    gamedbs = load_json(a.gamedbs)
    images = load_json(os.path.join(ROOT, "sprites", "registry.json"))["images"]

    wb = openpyxl.load_workbook(a.xlsx)
    ws = wb.active
    rows = read_rows(ws)
    report = dict(removed_no_commu=[], kept_no_commu_with_notes=[], added=[], excluded_no_commu=[], coll_changed=[], links_added=0,
                  new_events=[], unlisted_pickups=[], no_image=[])

    # ---- 0. コミュのないサポートカードを外す ------------------------------------
    for x in reversed(rows):
        if x["kind"] != S or not x["tags"]:
            continue
        w = best_match(x["name"], [c for c in wiki["cards"] if c["kind"] == "S" and c["idol"] == x["tags"][0]], lambda c: c["name"])
        if not w or w.get("commu") != 0:
            continue
        label = f'{iso(x["date"])} 【{x["name"]}】{x["tags"][0]}'
        if ws.cell(x["r"], COL["must"]).value not in (None, "") or ws.cell(x["r"], COL["memo"]).value not in (None, ""):
            report["kept_no_commu_with_notes"].append(label)
            continue
        delete_row(ws, x["r"])
        report["removed_no_commu"].append(label)
    rows = read_rows(ws)

    # ---- 1. 新しいカード ----------------------------------------------------
    card_dates = [x["date"] for x in rows if x["kind"] in CARD_KINDS and x["date"]]
    since = max(card_dates) - datetime.timedelta(days=NEW_CARD_WINDOW_DAYS)
    style_row = next(x["r"] for x in rows if x["kind"] in CARD_KINDS)
    new_cards = []
    for c in wiki["cards"]:
        d = to_date(c["date"])
        kind = KIND_OF[c["kind"]]
        if d < since or find_card(rows, kind, c["idol"], c["name"]):
            continue
        if kind == S and not c.get("commu"):
            report["excluded_no_commu"].append(f'{c["date"]} 【{c["name"]}】{c["idol"]}')
            continue
        new_cards.append(dict(c, d=d, kindname=kind))
    for c in sorted(new_cards, key=lambda c: c["d"]):  # 古い順に入れると、新しいものが上に来る
        at = next((x for x in read_rows(ws) if x["date"] and x["date"] <= c["d"]), None)
        at = at["r"] if at else ws.max_row + 1
        insert_rows(ws, at, 1, style_from=style_row if style_row < at else style_row + 1)
        style_row = style_row + 1 if style_row >= at else style_row
        vals = dict(tags=c["idol"], kind=c["kindname"], rarity=c["rarity"], name=c["name"],
                    date=datetime.datetime.combine(c["d"], datetime.time()))
        for k, v in vals.items():
            ws.cell(at, COL[k]).value = v
        report["added"].append(f'{c["date"]} {c["kindname"][:-4]} {c["rarity"]} 【{c["name"]}】{c["idol"]}')
    rows = read_rows(ws)

    # ---- 2. 区分 / 3. カードページ ---------------------------------------------
    for x in rows:
        if x["kind"] not in CARD_KINDS or not x["tags"]:
            continue
        idol, kc = x["tags"][0], "P" if x["kind"] == P else "S"
        w = best_match(x["name"], [c for c in wiki["cards"] if c["kind"] == kc and c["idol"] == idol], lambda c: c["name"])
        b = best_match(x["name"], [c for c in brilliant if c["kind"] == kc and idol in c["idols"]], lambda c: c["name"])
        coll = GACHA_TO_COLL.get(w["gacha"]) if w else None
        if not coll and b:
            extra = [t for t in b["colls"] if t in BRILLIANT_ONLY_COLLS]
            coll = "、".join(extra) if extra else None
        if coll and coll != x["coll"]:
            report["coll_changed"].append(f'【{x["name"]}】{idol}: {x["coll"] or "なし"} → {coll}')
            ws.cell(x["r"], COL["coll"]).value = coll
        if not x["link"]:
            url = b["url"] if b else None
            if not url:
                g = best_match(x["name"], [c for c in gamedbs if c["idol"] == idol], lambda c: c["name"])
                url = g["url"] if g else None
            if url:
                set_link(ws, x["r"], url)
                report["links_added"] += 1
    rows = read_rows(ws)

    # ---- 4. 画像のない カード → スプライト計画 --------------------------------
    by_url = {c["url"]: c for c in brilliant}
    plan = {"brilliant": [], "gamedbs": []}
    for x in rows:
        if x["kind"] not in CARD_KINDS:
            continue
        if not x["link"]:
            report["no_image"].append(f'【{x["name"]}】{"・".join(x["tags"])}（カードページなし）')
        elif not image_of(images, x["link"]):
            if x["link"] in by_url and by_url[x["link"]]["thumb"]:
                plan["brilliant"].append({"link": x["link"], "img": by_url[x["link"]]["thumb"]})
            elif "gamedbs.jp" in x["link"]:
                plan["gamedbs"].append({"link": x["link"], "page": x["link"]})

    # ---- 5. 総覧にないイベント / 一覧に未反映のピックアップ -------------------------
    events = [x for x in rows if x["kind"] == EVENT and x["date"]]
    idols = {t for x in rows if x["kind"] in CARD_KINDS for t in x["tags"]}
    last_event = max(x["date"] for x in events)
    for c in wiki["cards"]:
        if c["kind"] != "S" or c.get("get") != "シナリオ" or to_date(c["date"]) < last_event - datetime.timedelta(days=60):
            continue
        d = to_date(c["date"])
        if not any(abs((e["date"] - d).days) <= 3 and c["idol"] in people_of(e["tags"], idols) for e in events):
            report["new_events"].append(f'{c["date"]} 報酬【{c["name"]}】{c["idol"]}')
    listed = {(norm(c["name"]), c["idol"]) for c in wiki["cards"]}
    known = sorted({c["idol"] for c in wiki["cards"]}, key=len, reverse=True)
    for p in wiki.get("pickups", []):
        # ガシャ表のセルは説明文が続けて入っていることがあるので、既知のアイドル名で切り出す
        idol = next((i for i in known if p["idol"].startswith(i)), None)
        if not idol or to_date(p["start"]) < since or (norm(p["name"]), idol) in listed:
            continue
        s = f'{p["start"]} 【{p["name"]}】{idol}'
        if s not in report["unlisted_pickups"]:
            report["unlisted_pickups"].append(s)

    wb.save(a.out)
    dump_json(plan, a.plan)
    report["sprite_plan"] = {k: len(v) for k, v in plan.items()}
    dump_json(report, a.report)
    for k, v in report.items():
        print(f"{k}: {v}")


if __name__ == "__main__":
    main()
