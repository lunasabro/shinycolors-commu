"""総覧の Excel から公開サイト用の data.json を作る。

使い方:
  python scripts/build_data.py 総覧.xlsx --wiki work/wiki.json --brilliant work/brilliant.json \
      [--support-reflect 2026-10-11 --produce-reflect 2026-09-29] [-o data.json]

画像（img）の決め方:
  - カード行   : I列のカードページURL → sprites/registry.json
  - イベント行 : overrides.json → wiki のイベント報酬SSR（追加日±3日・登場人物）→ そのカードの画像
                 → 見つからなければ前回の data.json の値
  - その他の行 : 前回の data.json の値。新しい行なら、同じ実装日の登場人物のカード（サポート→SSR優先）

出力しないもの: G列のメモ（作成者の覚え書き）、I列のカードページURL（画像を引くためだけに使う）
"""
import argparse
import datetime
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from common import (CARD_KINDS, COL, COLLECTIONS, EVENT, S, UNITS, best_match, dump_json, image_of, iso, load_json,
                    people_of, split_tags, to_date)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# data.json の各行の並び。サイト（assets/app.js）は data.json の "cols" を見て読むので、ここだけ直せばよい
COLS = ["tags", "kind", "rarity", "name", "date", "must", "coll", "img"]


def read_rows(xlsx):
    ws = openpyxl.load_workbook(xlsx).active
    rows = []
    for r in range(2, ws.max_row + 1):
        v = {k: ws.cell(r, c).value for k, c in COL.items()}
        if not any([v["tags"], v["kind"], v["name"]]):
            continue
        cell = ws.cell(r, COL["link"])
        rows.append(dict(
            tags=split_tags(v["tags"]), kind=v["kind"] or "", rarity=v["rarity"] or "",
            name="" if v["name"] is None else str(v["name"]), date=iso(v["date"]),
            must=int(v["must"]) if isinstance(v["must"], (int, float)) else None,
            coll=[c for c in (v["coll"] or "").split("、") if c],
            link=cell.hyperlink.target if cell.hyperlink else "",
        ))
    return rows


def key(r):
    return (r["kind"], tuple(r["tags"]), r["name"], r["date"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--wiki", help="scrape_wiki.js の結果（イベント報酬の照合と反映日に使う）")
    ap.add_argument("--brilliant", help="scrape_brilliant.js の結果")
    ap.add_argument("--support-reflect")
    ap.add_argument("--produce-reflect")
    ap.add_argument("-o", "--out", default=os.path.join(ROOT, "data.json"))
    ap.add_argument("--prev", default=os.path.join(ROOT, "data.json"))
    ap.add_argument("--today", default=datetime.date.today().isoformat())
    a = ap.parse_args()

    reg = load_json(os.path.join(ROOT, "sprites", "registry.json"))
    images = reg["images"]
    overrides = load_json(os.path.join(ROOT, "scripts", "overrides.json"), {}).get("event_images", {})
    prev = load_json(a.prev, {"meta": {}, "rows": []})
    pc = {c: i for i, c in enumerate(prev.get("cols", COLS))}
    prev_img = {(p[pc["kind"]], tuple(p[pc["tags"]]), p[pc["name"]], p[pc["date"]]): p[pc["img"]]
                for p in prev["rows"]}
    wiki = load_json(a.wiki, {}) if a.wiki else {}
    brilliant = load_json(a.brilliant, []) if a.brilliant else []

    rows = read_rows(a.xlsx)
    cards = [r for r in rows if r["kind"] in CARD_KINDS]
    idols = {t for r in cards for t in r["tags"]}

    def card_img(r):
        return image_of(images, r["link"])

    def img_of_card(name, idol):
        """カード名＋アイドル名から画像を探す（総覧の行 → BRILLI@NT 一覧）。"""
        c = best_match(name, [r for r in cards if idol in r["tags"]], lambda r: r["name"])
        if c and card_img(c):
            return card_img(c)
        b = best_match(name, [x for x in brilliant if idol in x["idols"]], lambda x: x["name"], 0.85)
        return image_of(images, b["url"]) if b else None

    # イベント報酬：サポート一覧で「入手」が「シナリオ」のカード
    rewards = [dict(name=c["name"], idol=c["idol"], rarity=c["rarity"], d=to_date(c["date"]))
               for c in wiki.get("cards", []) if c["kind"] == "S" and c.get("get") == "シナリオ"]
    stats = dict(event_from_reward=0, event_kept=0, event_none=[])

    out = []
    for r in rows:
        k = key(r)
        if r["kind"] in CARD_KINDS:
            img = card_img(r) or prev_img.get(k)
        elif r["kind"] == EVENT:
            img = None
            ov = overrides.get(f'{r["kind"]}|{r["name"]}|{r["date"]}')
            if ov:
                img = image_of(images, ov)
            elif rewards and r["date"]:
                d, ppl = to_date(r["date"]), people_of(r["tags"], idols)
                near = [x for x in rewards if abs((x["d"] - d).days) <= 3]
                cand = [x for x in near if x["idol"] in ppl] or ([] if ppl else near)
                if cand:
                    x = min(cand, key=lambda x: (abs((x["d"] - d).days), x["rarity"] != "SSR"))
                    img = img_of_card(x["name"], x["idol"])
            if img:
                stats["event_from_reward"] += 1
            elif k in prev_img and prev_img[k]:
                img = prev_img[k]
                stats["event_kept"] += 1
            else:
                stats["event_none"].append(r["name"])
        else:
            if k in prev_img:
                img = prev_img[k]
            else:
                ppl = people_of(r["tags"], idols)
                same = [c for c in cards if c["date"] == r["date"] and set(c["tags"]) & ppl and card_img(c)]
                same.sort(key=lambda c: (c["kind"] != S, c["rarity"] != "SSR"))
                img = card_img(same[0]) if same else None
        r["img"] = img
        out.append([r[c] for c in COLS])

    meta = dict(prev.get("meta", {}))
    meta["updated"] = a.today
    for k, w in (("support_wiki", "support_reflect"), ("produce_wiki", "produce_reflect")):
        v = getattr(a, w) or wiki.get(w)
        if v:
            meta[k] = v
    meta["sheets"] = reg["sheets"]
    meta["units"] = UNITS
    meta["collections"] = COLLECTIONS
    data = {"meta": meta, "cols": COLS, "rows": out}
    dump_json(data, a.out, compact=True)
    print(f"rows={len(out)} images={sum(1 for r in rows if r["img"])} "
          f"event_from_reward={stats['event_from_reward']} event_kept={stats['event_kept']} "
          f"event_none={stats['event_none']}")


if __name__ == "__main__":
    main()
